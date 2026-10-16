"""Create and inspect editable Excel workbooks without overwriting existing files."""
from __future__ import annotations

from core.path_policy import PathPolicyError, resolve_user_path, user_accessible_roots
from actions._office_common import safe_filename, safe_output_directory, unique_output_path


def office_spreadsheets(parameters: dict, player=None, speak=None) -> str:
    params = parameters or {}
    operation = str(params.get("operation", "create_workbook")).strip().lower()
    try:
        from openpyxl import Workbook, load_workbook
    except ImportError:
        return "openpyxl is not installed. Install it with: pip install openpyxl"

    try:
        if operation == "create_workbook":
            title = str(params.get("title", "Workbook")).strip() or "Workbook"
            headers = params.get("headers", [])
            rows = params.get("rows", [])
            if not isinstance(headers, list) or any(not isinstance(v, (str, int, float, bool, type(None))) for v in headers):
                return "headers must be a list of simple cell values."
            if not isinstance(rows, list) or any(not isinstance(row, list) for row in rows):
                return "rows must be a list of row lists."
            if any(len(row) > 500 for row in rows) or len(rows) > 10000:
                return "Workbook input exceeds the limit of 10,000 rows or 500 columns per row."
            if any(not isinstance(v, (str, int, float, bool, type(None))) for row in rows for v in row):
                return "Spreadsheet cells must be strings, numbers, booleans, or null."
            directory = safe_output_directory(str(params.get("output_dir", "") or ""))
            output = unique_output_path(directory, safe_filename(str(params.get("filename", "") or ""), title, ".xlsx"))
            workbook = Workbook()
            sheet_name = str(params.get("sheet_name", "Sheet1") or "Sheet1").replace("[", "(").replace("]", ")")
            sheet_name = sheet_name[:31] or "Sheet1"
            sheet = workbook.active
            sheet.title = sheet_name
            sheet.freeze_panes = "A2" if headers else None
            if headers:
                sheet.append(headers)
                for cell in sheet[1]:
                    cell.font = cell.font.copy(bold=True)
                    cell.data_type = "s" if isinstance(cell.value, str) else cell.data_type
            for row in rows:
                sheet.append(row)
            # Treat all user-provided strings as text, not executable formulas.
            for row in sheet.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str):
                        cell.data_type = "s"
            workbook.save(output)
            with load_workbook(output, read_only=True, data_only=False) as check:
                if sheet_name not in check.sheetnames:
                    output.unlink(missing_ok=True)
                    return "Workbook validation failed; incomplete output removed."
                checked_rows = check[sheet_name].max_row
                checked_cols = check[sheet_name].max_column
            return f"Created editable XLSX: {output}\nSheet: {sheet_name}\nRows: {checked_rows}; columns: {checked_cols}"

        if operation == "inspect_workbook":
            raw = str(params.get("file_path", "") or "").strip()
            if not raw:
                return "file_path is required for inspect_workbook."
            source = resolve_user_path(raw, roots=user_accessible_roots(), must_exist=True)
            if not source.is_file() or source.suffix.lower() != ".xlsx":
                return "Please provide an existing .xlsx file."
            with load_workbook(source, read_only=True, data_only=True) as workbook:
                parts = [f"Workbook: {source}", f"Sheets: {', '.join(workbook.sheetnames)}"]
                for sheet in workbook.worksheets[:10]:
                    parts.append(f"\n[{sheet.title}] {sheet.max_row} rows × {sheet.max_column} columns")
                    for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 8), values_only=True):
                        parts.append(" | ".join("" if value is None else str(value)[:200] for value in row[:20]))
                return "\n".join(parts)[:12000]
        return "Unsupported operation. Use create_workbook or inspect_workbook."
    except (OSError, ValueError, TypeError, PathPolicyError) as exc:
        return f"Spreadsheet operation failed safely: {exc}"
    except Exception as exc:
        return f"Spreadsheet operation failed: {exc}"


TOOL = {
    "name": "office_spreadsheets",
    "description": "Create editable Excel XLSX workbooks from headers and rows, or inspect an existing XLSX. Input strings are written as text, not formulas. Existing files are never overwritten.",
    "parameters": {
        "type": "OBJECT",
        "properties": {
            "operation": {"type": "STRING", "description": "create_workbook | inspect_workbook"},
            "title": {"type": "STRING", "description": "Workbook title and default filename"},
            "filename": {"type": "STRING", "description": "Optional filename only"},
            "output_dir": {"type": "STRING", "description": "Optional output folder inside the current user's home directory or OS temporary directory"},
            "sheet_name": {"type": "STRING", "description": "Optional worksheet name"},
            "headers": {"type": "ARRAY", "items": {"type": "STRING"}, "description": "Optional first row of column headers"},
            "rows": {"type": "ARRAY", "items": {"type": "ARRAY", "items": {"type": "STRING"}}, "description": "Rows of text, numbers, booleans or null values"},
            "file_path": {"type": "STRING", "description": "Existing XLSX path for inspect_workbook"}
        },
        "required": ["operation"]
    },
    "handler": office_spreadsheets
}
