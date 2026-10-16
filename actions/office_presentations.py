"""Create editable PowerPoint presentations locally."""
from __future__ import annotations

from core.path_policy import PathPolicyError
from actions._office_common import safe_filename, safe_output_directory, unique_output_path


def office_presentations(parameters: dict, player=None, speak=None) -> str:
    params = parameters or {}
    operation = str(params.get("operation", "create_presentation")).strip().lower()
    if operation != "create_presentation":
        return "Unsupported operation. Use create_presentation."
    title = str(params.get("title", "Presentation")).strip() or "Presentation"
    slides = params.get("slides", [])
    if not isinstance(slides, list) or not slides:
        return "slides must be a non-empty list of slide objects with title and optional bullets."
    if len(slides) > 100:
        return "A presentation is limited to 100 slides per operation."
    if any(not isinstance(slide, dict) for slide in slides):
        return "Every slide must be an object with a title and optional bullets."
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
        from pptx.enum.text import PP_ALIGN

        directory = safe_output_directory(str(params.get("output_dir", "") or ""))
        output = unique_output_path(directory, safe_filename(str(params.get("filename", "") or ""), title, ".pptx"))
        presentation = Presentation()
        presentation.core_properties.title = title
        presentation.core_properties.subject = "Created by Mark-LV Office Productivity Agent"
        presentation.slide_width = Inches(13.333)
        presentation.slide_height = Inches(7.5)

        for index, item in enumerate(slides):
            slide_title = str(item.get("title", "") or (title if index == 0 else f"Slide {index + 1}")).strip()
            bullets = item.get("bullets", [])
            body = item.get("body", "")
            if not isinstance(bullets, list) or any(not isinstance(v, str) for v in bullets):
                return f"Slide {index + 1}: bullets must be a list of strings."
            if not isinstance(body, str):
                return f"Slide {index + 1}: body must be text."
            layout = presentation.slide_layouts[0] if index == 0 else presentation.slide_layouts[1]
            slide = presentation.slides.add_slide(layout)
            slide.shapes.title.text = slide_title
            if index == 0:
                subtitle = slide.placeholders[1] if len(slide.placeholders) > 1 else None
                if subtitle is not None:
                    subtitle.text = body or "Prepared with Mark-LV"
            else:
                placeholder = slide.placeholders[1] if len(slide.placeholders) > 1 else None
                if placeholder is not None:
                    frame = placeholder.text_frame
                    frame.clear()
                    if body.strip():
                        paragraph = frame.paragraphs[0]
                        paragraph.text = body.strip()
                        paragraph.font.size = Pt(18)
                    for bullet in bullets:
                        paragraph = frame.add_paragraph()
                        paragraph.text = bullet
                        paragraph.level = 0
                        paragraph.font.size = Pt(20)
        presentation.save(output)
        check = Presentation(output)
        if len(check.slides) != len(slides):
            output.unlink(missing_ok=True)
            return "Presentation validation failed; incomplete output removed."
        return f"Created editable PPTX: {output}\nSlides: {len(check.slides)}"
    except ImportError:
        return "python-pptx is not installed. Install it with: pip install python-pptx"
    except (OSError, ValueError, TypeError, PathPolicyError) as exc:
        return f"Presentation operation failed safely: {exc}"
    except Exception as exc:
        return f"Presentation operation failed: {exc}"


TOOL = {
    "name": "office_presentations",
    "description": "Create editable PowerPoint PPTX presentations from slide titles, body text and bullet lists. Existing files are never overwritten.",
    "parameters": {
        "type": "OBJECT",
        "properties": {
            "operation": {"type": "STRING", "description": "create_presentation"},
            "title": {"type": "STRING", "description": "Presentation title and default filename"},
            "filename": {"type": "STRING", "description": "Optional filename only"},
            "output_dir": {"type": "STRING", "description": "Optional output folder inside the current user's home directory or OS temporary directory"},
            "slides": {"type": "ARRAY", "items": {"type": "OBJECT"}, "description": "Slide objects with title, optional body, and optional bullets array"}
        },
        "required": ["operation", "title", "slides"]
    },
    "handler": office_presentations
}
