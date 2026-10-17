"""Generate/edit visual assets with Gemini and optionally insert them into Office files."""
from __future__ import annotations

import asyncio
import base64
import json
import mimetypes
import os
import posixpath
import sys
import uuid
import zipfile
import html
from pathlib import Path
from xml.etree import ElementTree as ET

from actions._office_common import safe_filename, safe_output_directory, unique_output_path
from core.path_policy import PathPolicyError, resolve_user_path, user_accessible_roots

_MAX_PROMPT = 12000
_MAX_SOURCE_BYTES = 20 * 1024 * 1024
_MAX_OUTPUT_BYTES = 30 * 1024 * 1024
_MAX_PIXELS = 100_000_000
_MODEL_DEFAULT = "gemini-2.5-flash-image"
_ALLOWED_ASPECT_RATIOS = {"1:1", "2:3", "3:2", "3:4", "4:3", "4:5", "5:4", "9:16", "16:9", "21:9"}
_ALLOWED_IMAGE_SIZES = {"1K", "2K", "4K"}
_IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp"}

_MAX_EPUB_ENTRIES = 10000
_MAX_EPUB_MEMBER_BYTES = 100 * 1024 * 1024
_MAX_EPUB_EXPANDED_BYTES = 750 * 1024 * 1024
_MAX_EPUB_COMPRESSION_RATIO = 1000


def _validate_epub_archive(path: Path) -> tuple[list[str], str]:
    """Validate EPUB ZIP structure and resource bounds without extracting members."""
    if not zipfile.is_zipfile(path):
        raise ValueError("Target EPUB is not a valid ZIP package.")
    try:
        with zipfile.ZipFile(path, "r") as archive:
            infos = archive.infolist()
            if not infos or len(infos) > _MAX_EPUB_ENTRIES:
                raise ValueError(f"EPUB must contain between 1 and {_MAX_EPUB_ENTRIES} archive entries.")
            names = [info.filename for info in infos]
            if len(names) != len(set(names)):
                raise ValueError("EPUB contains duplicate archive paths; refusing ambiguous edits.")
            if names[0] != "mimetype" or infos[0].compress_type != zipfile.ZIP_STORED:
                raise ValueError("EPUB must begin with an uncompressed mimetype entry.")
            if infos[0].file_size != len(b"application/epub+zip"):
                raise ValueError("EPUB mimetype entry is missing or invalid.")
            if archive.read("mimetype") != b"application/epub+zip":
                raise ValueError("EPUB mimetype entry is missing or invalid.")

            expanded_total = 0
            for info in infos:
                name = info.filename
                normalized = posixpath.normpath(name)
                if (
                    not name or name.startswith("/") or "\\" in name
                    or normalized in {".", ".."} or normalized.startswith("../")
                    or (len(name) >= 2 and name[1] == ":")
                ):
                    raise ValueError("EPUB contains an unsafe archive path.")
                mode = (info.external_attr >> 16) & 0o170000
                if mode == 0o120000:
                    raise ValueError("EPUB symbolic links are not supported.")
                if info.file_size > _MAX_EPUB_MEMBER_BYTES:
                    raise ValueError("An EPUB archive member exceeds the 100 MiB safety limit.")
                expanded_total += info.file_size
                if expanded_total > _MAX_EPUB_EXPANDED_BYTES:
                    raise ValueError("EPUB expanded contents exceed the 750 MiB safety limit.")
                if info.file_size > 1024 * 1024 and info.compress_size > 0:
                    if info.file_size / info.compress_size > _MAX_EPUB_COMPRESSION_RATIO:
                        raise ValueError("EPUB contains a suspiciously high compression ratio.")

            container = ET.fromstring(archive.read("META-INF/container.xml"))
            rootfile = next(
                (node for node in container.iter() if node.tag.rsplit("}", 1)[-1] == "rootfile"),
                None,
            )
            opf_path = rootfile.attrib.get("full-path", "") if rootfile is not None else ""
            normalized_opf = posixpath.normpath(opf_path)
            if (
                not opf_path or normalized_opf != opf_path or opf_path not in names
                or opf_path.startswith("/") or normalized_opf == ".."
                or normalized_opf.startswith("../")
            ):
                raise ValueError("EPUB package document could not be located safely.")
            ET.fromstring(archive.read(opf_path))
            return names, opf_path
    except (KeyError, ET.ParseError, zipfile.BadZipFile, RuntimeError) as exc:
        raise ValueError("EPUB package is incomplete, encrypted or malformed.") from exc



def _api_key() -> str:
    """Read the same local API-key configuration used by the main Gemini client."""
    base_dir = Path(sys.executable).resolve().parent if getattr(sys, "frozen", False) else Path(__file__).resolve().parent.parent
    config_path = base_dir / "config" / "api_keys.json"
    try:
        data = json.loads(config_path.read_text(encoding="utf-8"))
        key = str(data.get("gemini_api_key", "")).strip()
    except (OSError, ValueError, TypeError):
        key = ""
    if not key:
        key = os.environ.get("GEMINI_API_KEY", "").strip()
    if not key:
        raise ValueError("Gemini API key is not configured. Add gemini_api_key to config/api_keys.json.")
    return key


def _image_path(raw: str) -> Path:
    if not isinstance(raw, str) or not raw.strip():
        raise ValueError("image_path is required for image editing.")
    path = resolve_user_path(raw.strip(), roots=user_accessible_roots(), must_exist=True)
    if not path.is_file() or path.suffix.lower() not in _IMAGE_EXTENSIONS:
        raise ValueError("image_path must point to an existing PNG, JPEG or WEBP image.")
    if path.stat().st_size > _MAX_SOURCE_BYTES:
        raise ValueError("Source image exceeds the 20 MiB limit.")
    from PIL import Image
    with Image.open(path) as image:
        image.verify()
    with Image.open(path) as image:
        if image.width <= 0 or image.height <= 0 or image.width * image.height > _MAX_PIXELS:
            raise ValueError("Source image exceeds the safe pixel limit.")
        image.load()
    return path


def _extract_image(response) -> tuple[bytes, str]:
    """Extract only actual image parts; never treat model text as an image result."""
    for candidate in getattr(response, "candidates", None) or []:
        content = getattr(candidate, "content", None)
        for part in getattr(content, "parts", None) or []:
            inline = getattr(part, "inline_data", None)
            data = getattr(inline, "data", None) if inline else None
            mime = str(getattr(inline, "mime_type", "") or "image/png").lower() if inline else ""
            if data and mime in {"image/png", "image/jpeg", "image/webp"}:
                if isinstance(data, str):
                    import base64
                    data = base64.b64decode(data, validate=True)
                if len(data) <= _MAX_OUTPUT_BYTES:
                    return bytes(data), mime
    raise ValueError("The image model returned no usable image. Try a clearer prompt or another model.")


def _font(size: int):
    """Choose a dependable local font without requiring font installation."""
    from PIL import ImageFont
    candidates = [
        Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts" / "arial.ttf",
        Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts" / "segoeui.ttf",
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
        Path("/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf"),
    ]
    for candidate in candidates:
        try:
            if candidate.is_file():
                return ImageFont.truetype(str(candidate), size=size)
        except OSError:
            continue
    return ImageFont.load_default()


def _iter_docx_paragraphs(container, seen=None):
    """Yield document paragraphs, including nested table cells, without duplicates."""
    if seen is None:
        seen = set()
    for paragraph in getattr(container, "paragraphs", ()):
        identity = id(paragraph._p)
        if identity not in seen:
            seen.add(identity)
            yield paragraph
    for table in getattr(container, "tables", ()):
        for row in table.rows:
            for cell in row.cells:
                yield from _iter_docx_paragraphs(cell, seen)


def _compose_cover_text(image_path: Path, params: dict) -> Path:
    """Add readable, deterministic cover typography while retaining the raw asset."""
    title = str(params.get("cover_title", "") or "").strip()
    subtitle = str(params.get("cover_subtitle", "") or "").strip()
    author = str(params.get("cover_author", "") or "").strip()
    if not any((title, subtitle, author)):
        return image_path
    from PIL import Image, ImageDraw
    with Image.open(image_path) as opened:
        opened.load()
        if opened.width * opened.height > _MAX_PIXELS:
            raise ValueError("Generated image exceeds the safe pixel limit.")
        base = opened.convert("RGBA")
    width, height = base.size
    overlay = Image.new("RGBA", base.size, (0, 0, 0, 0))
    gradient_draw = ImageDraw.Draw(overlay)
    start_y = int(height * 0.52)
    for y in range(start_y, height):
        alpha = int(35 + 175 * ((y - start_y) / max(1, height - start_y)))
        gradient_draw.line((0, y, width, y), fill=(8, 12, 20, alpha))
    composed = Image.alpha_composite(base, overlay)
    draw = ImageDraw.Draw(composed)
    margin = max(28, int(width * 0.075))
    available = max(100, width - margin * 2)
    title_font = _font(max(20, min(84, int(width * 0.075))))
    subtitle_font = _font(max(14, min(42, int(width * 0.035))))
    author_font = _font(max(14, min(36, int(width * 0.028))))

    def wrap(text: str, font, max_width: int):
        lines, current = [], ""
        for word in text.split():
            candidate = f"{current} {word}".strip()
            if current and draw.textbbox((0, 0), candidate, font=font)[2] > max_width:
                lines.append(current)
                current = word
            else:
                current = candidate
        if current:
            lines.append(current)
        return lines

    blocks = []
    if title:
        blocks.append((title_font, wrap(title, title_font, available), (255, 255, 255, 255), max(4, int(height * 0.012))))
    if subtitle:
        blocks.append((subtitle_font, wrap(subtitle, subtitle_font, available), (235, 239, 245, 255), max(3, int(height * 0.008))))
    if author:
        blocks.append((author_font, wrap(author, author_font, available), (240, 196, 90, 255), max(4, int(height * 0.018))))
    line_gap = max(4, int(height * 0.008))
    measured, total = [], 0
    for font, lines, fill, gap in blocks:
        heights = [draw.textbbox((0, 0), line, font=font)[3] - draw.textbbox((0, 0), line, font=font)[1] for line in lines]
        block_height = sum(heights) + line_gap * max(0, len(lines) - 1)
        measured.append((font, lines, fill, gap, heights))
        total += block_height + gap
    y = max(start_y + int(height * 0.04), height - margin - total)
    for font, lines, fill, gap, heights in measured:
        for line, line_height in zip(lines, heights):
            draw.text((margin, y), line, font=font, fill=fill, stroke_width=max(1, width // 1200), stroke_fill=(0, 0, 0, 180))
            y += line_height + line_gap
        y += gap
    output = unique_output_path(image_path.parent, safe_filename(image_path.stem + "-designed", "designed-cover", image_path.suffix))
    temporary = output.with_name(f".{output.stem}-{uuid.uuid4().hex}.tmp{output.suffix}")
    try:
        if output.suffix.lower() in {".jpg", ".jpeg"}:
            composed.convert("RGB").save(temporary, format="JPEG", quality=94, optimize=True)
        elif output.suffix.lower() == ".webp":
            composed.save(temporary, format="WEBP", quality=94, method=6)
        else:
            composed.save(temporary, format="PNG", optimize=True)
        with Image.open(temporary) as check:
            check.verify()
        os.replace(temporary, output)
    finally:
        temporary.unlink(missing_ok=True)
    return output


def _preflight_target(raw_target: str, slide_index=None, pptx_placement=None, shape_index=None, docx_insertion_mode=None, insert_after_text=None) -> Path | None:
    """Validate an Office target before consuming image-generation quota."""
    if not raw_target.strip():
        return None
    target = resolve_user_path(raw_target.strip(), roots=user_accessible_roots(), must_exist=True)
    if not target.is_file() or target.suffix.lower() not in {".docx", ".pptx", ".epub"}:
        raise ValueError("target_file_path must point to an existing .docx, .pptx or .epub file.")
    if target.stat().st_size > 100 * 1024 * 1024:
        raise ValueError("Target document exceeds the 100 MiB limit.")
    if target.suffix.lower() == ".epub":
        _validate_epub_archive(target)
    if target.suffix.lower() == ".pptx" and slide_index is not None:
        from pptx import Presentation
        presentation = Presentation(target)
        if isinstance(slide_index, bool) or not isinstance(slide_index, int) or not 0 <= slide_index < len(presentation.slides):
            raise ValueError(f"slide_index must be between 0 and {len(presentation.slides) - 1}.")
        if str(pptx_placement or "").lower() == "replace_picture":
            if isinstance(shape_index, bool) or not isinstance(shape_index, int):
                raise ValueError("shape_index is required for pptx_placement=replace_picture.")
            shapes = presentation.slides[slide_index].shapes
            if not 0 <= shape_index < len(shapes):
                raise ValueError(f"shape_index must be between 0 and {len(shapes) - 1}.")
            if shapes[shape_index].shape_type != 13:
                raise ValueError("replace_picture requires shape_index to identify an existing picture.")
    elif str(pptx_placement or "").lower() == "replace_picture":
        raise ValueError("replace_picture requires slide_index and shape_index.")
    if target.suffix.lower() == ".docx" and str(docx_insertion_mode or "").lower() == "inline" and str(insert_after_text or "").strip():
        from docx import Document
        document = Document(target)
        anchor = str(insert_after_text).strip()
        if not any(anchor in paragraph.text for paragraph in _iter_docx_paragraphs(document)):
            raise ValueError(f"insert_after_text was not found in the document: {anchor[:120]}")
    return target


def _generate(params: dict) -> tuple[Path, str]:
    prompt = str(params.get("prompt", "") or "").strip()
    if not prompt:
        raise ValueError("prompt is required.")
    if len(prompt) > _MAX_PROMPT:
        raise ValueError("prompt must not exceed 12,000 characters.")
    operation = str(params.get("operation", "generate")).strip().lower()
    if operation not in {"generate", "edit"}:
        raise ValueError("operation must be generate or edit.")
    source = _image_path(str(params.get("image_path", "") or "")) if operation == "edit" else None

    # Keep the existing application key store as the source of credentials.
    from google import genai
    from google.genai import types
    client = genai.Client(api_key=_api_key())
    contents = [prompt]
    if source:
        source_mime = {
            ".png": "image/png", ".jpg": "image/jpeg",
            ".jpeg": "image/jpeg", ".webp": "image/webp",
        }[source.suffix.lower()]
        contents.append(types.Part.from_bytes(data=source.read_bytes(), mime_type=source_mime))
    model = str(params.get("model", "") or _MODEL_DEFAULT).strip()
    if not model or len(model) > 100 or any(ch.isspace() for ch in model):
        raise ValueError("model must be a valid model identifier.")
    aspect_ratio = str(params.get("aspect_ratio", "") or "").strip()
    image_size = str(params.get("image_size", "") or "").strip().upper()
    if aspect_ratio and aspect_ratio not in _ALLOWED_ASPECT_RATIOS:
        raise ValueError("Unsupported aspect_ratio. Use a documented image ratio.")
    if image_size and image_size not in _ALLOWED_IMAGE_SIZES:
        raise ValueError("image_size must be 1K, 2K or 4K.")
    config_options = {"response_modalities": ["TEXT", "IMAGE"]}
    if aspect_ratio or image_size:
        image_config = {}
        if aspect_ratio:
            image_config["aspect_ratio"] = aspect_ratio
        if image_size:
            image_config["image_size"] = image_size
        config_options["image_config"] = types.ImageConfig(**image_config)
    response = client.models.generate_content(
        model=model,
        contents=contents,
        config=types.GenerateContentConfig(**config_options),
    )
    image_bytes, mime = _extract_image(response)
    extension = { "image/png": ".png", "image/jpeg": ".jpg", "image/webp": ".webp" }[mime]
    directory_raw = str(params.get("output_dir", "") or "")
    directory = safe_output_directory(directory_raw) if directory_raw else safe_output_directory()
    default_name = str(params.get("filename", "") or ("visual-edit" if source else "generated-visual"))
    filename = safe_filename(default_name, "generated-visual", extension)
    # Always align the extension with the actual MIME type returned by the model.
    filename = str(Path(filename).with_suffix(extension))
    output = unique_output_path(directory, filename)
    temporary = output.with_name(f".{output.stem}-{uuid.uuid4().hex}.tmp{extension}")
    try:
        temporary.write_bytes(image_bytes)
        from PIL import Image
        with Image.open(temporary) as image:
            image.verify()
        with Image.open(temporary) as image:
            if image.width * image.height > _MAX_PIXELS:
                raise ValueError("Generated image exceeds the safe pixel limit.")
            image.load()
    except Exception:
        temporary.unlink(missing_ok=True)
        raise
    os.replace(temporary, output)
    return _compose_cover_text(output, params), mime


def _insert_docx(image_path: Path, raw_target: str, insertion_mode: str = "cover", caption: str = "", insert_after_text: str = "", alt_text: str = "") -> Path:
    from docx import Document
    from docx.enum.text import WD_BREAK
    from docx.shared import Inches
    target = resolve_user_path(raw_target, roots=user_accessible_roots(), must_exist=True)
    if not target.is_file() or target.suffix.lower() != ".docx":
        raise ValueError("DOCX insertion requires an existing .docx file.")
    if target.stat().st_size > 100 * 1024 * 1024:
        raise ValueError("Target document exceeds the 100 MiB limit.")
    output = unique_output_path(target.parent, safe_filename(target.stem + "-visual", target.stem + "-visual", ".docx"))
    document = Document(target)
    if insertion_mode not in {"cover", "inline"}:
        raise ValueError("docx_insertion_mode must be cover or inline.")
    from PIL import Image
    with Image.open(image_path) as image:
        image_width, image_height = image.size
    max_width, max_height = Inches(5.8), Inches(7.5 if insertion_mode == "cover" else 6.0)
    scale = min(max_width / image_width, max_height / image_height)
    paragraph = document.add_paragraph()
    paragraph.alignment = 1
    inline_shape = paragraph.add_run().add_picture(str(image_path), width=int(image_width * scale), height=int(image_height * scale))
    if alt_text.strip():
        inline_shape._inline.docPr.set("descr", alt_text.strip()[:1000])
        inline_shape._inline.docPr.set("title", alt_text.strip()[:255])
    caption_paragraph = None
    if caption.strip():
        caption_paragraph = document.add_paragraph(caption.strip(), style="Caption")
        caption_paragraph.alignment = 1
    if insertion_mode == "inline" and insert_after_text.strip():
        anchor = next((p for p in _iter_docx_paragraphs(document) if insert_after_text.strip() in p.text), None)
        if anchor is None:
            raise ValueError(f"insert_after_text was not found in the document: {insert_after_text[:120]}")
        anchor._p.addnext(paragraph._p)
        if caption_paragraph is not None:
            paragraph._p.addnext(caption_paragraph._p)
    body = document._element.body
    if insertion_mode == "cover":
        page_break = document.add_paragraph()
        page_break.add_run().add_break(WD_BREAK.PAGE)
        body.insert(0, page_break._p)
        if caption_paragraph is not None:
            body.insert(0, caption_paragraph._p)
        body.insert(0, paragraph._p)
    temporary = output.with_name(f".{output.stem}-{uuid.uuid4().hex}.tmp.docx")
    try:
        document.save(temporary)
        with zipfile.ZipFile(temporary) as archive:
            if archive.testzip():
                raise ValueError("Generated Word document failed package validation.")
        os.replace(temporary, output)
    finally:
        temporary.unlink(missing_ok=True)
    return output


def _insert_pptx(image_path: Path, raw_target: str, slide_index=None, caption: str = "", placement: str = "append_slide", shape_index=None, alt_text: str = "") -> Path:
    from pptx import Presentation
    from pptx.util import Inches
    target = resolve_user_path(raw_target, roots=user_accessible_roots(), must_exist=True)
    if not target.is_file() or target.suffix.lower() != ".pptx":
        raise ValueError("PowerPoint insertion requires an existing .pptx file.")
    if target.stat().st_size > 100 * 1024 * 1024:
        raise ValueError("Target presentation exceeds the 100 MiB limit.")
    output = unique_output_path(target.parent, safe_filename(target.stem + "-visual", target.stem + "-visual", ".pptx"))
    presentation = Presentation(target)
    placement = str(placement or ("contain" if slide_index is not None else "append_slide")).strip().lower()
    if placement not in {"append_slide", "contain", "replace_picture"}:
        raise ValueError("pptx_placement must be append_slide, contain, or replace_picture.")
    if placement == "replace_picture" and (slide_index is None or shape_index is None):
        raise ValueError("replace_picture requires slide_index and shape_index.")
    from PIL import Image
    with Image.open(image_path) as image:
        image_width, image_height = image.size
    replaced_shape = None
    if placement == "append_slide":
        layout = next((item for item in presentation.slide_layouts if len(item.placeholders) == 0), presentation.slide_layouts[-1])
        slide = presentation.slides.add_slide(layout)
        bound_left = bound_top = Inches(0)
        bound_width, bound_height = presentation.slide_width, presentation.slide_height
    else:
        if isinstance(slide_index, bool) or not isinstance(slide_index, int) or not 0 <= slide_index < len(presentation.slides):
            raise ValueError(f"slide_index must be between 0 and {len(presentation.slides) - 1}.")
        slide = presentation.slides[slide_index]
        if placement == "replace_picture":
            if isinstance(shape_index, bool) or not isinstance(shape_index, int) or not 0 <= shape_index < len(slide.shapes):
                raise ValueError(f"shape_index must be between 0 and {len(slide.shapes) - 1}.")
            replaced_shape = slide.shapes[shape_index]
            if replaced_shape.shape_type != 13:
                raise ValueError("replace_picture requires shape_index to identify an existing picture.")
            bound_left, bound_top = replaced_shape.left, replaced_shape.top
            bound_width, bound_height = replaced_shape.width, replaced_shape.height
        else:
            bound_left = bound_top = Inches(1)
            bound_width = presentation.slide_width - Inches(2)
            bound_height = presentation.slide_height - Inches(2)
    scale = min(bound_width / image_width, bound_height / image_height)
    width, height = int(image_width * scale), int(image_height * scale)
    left = bound_left + int((bound_width - width) / 2)
    top = bound_top + int((bound_height - height) / 2)
    picture = slide.shapes.add_picture(str(image_path), left, top, width=width, height=height)
    picture.name = "Mark-LV generated visual"
    if alt_text.strip():
        picture._element.nvPicPr.cNvPr.set("descr", alt_text.strip()[:1000])
        picture._element.nvPicPr.cNvPr.set("title", alt_text.strip()[:255])
    if replaced_shape is not None:
        shape_tree = slide.shapes._spTree
        shape_tree.insert(shape_tree.index(replaced_shape._element), picture._element)
        shape_tree.remove(replaced_shape._element)
    if caption.strip():
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Pt
        box = slide.shapes.add_textbox(Inches(0.25), presentation.slide_height - Inches(0.48), presentation.slide_width - Inches(0.5), Inches(0.3))
        frame = box.text_frame
        frame.text = caption.strip()
        frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        frame.paragraphs[0].font.size = Pt(12)
    temporary = output.with_name(f".{output.stem}-{uuid.uuid4().hex}.tmp.pptx")
    try:
        presentation.save(temporary)
        with zipfile.ZipFile(temporary) as archive:
            if archive.testzip():
                raise ValueError("Generated PowerPoint failed package validation.")
        os.replace(temporary, output)
    finally:
        temporary.unlink(missing_ok=True)
    return output



def _insert_epub_cover(image_path: Path, raw_target: str, title: str = "", alt_text: str = "") -> Path:
    """Add an accessible cover image and cover page to a new EPUB copy.

    Existing archive members are retained; the original EPUB is never overwritten.
    Both EPUB 3 cover-image properties and the EPUB 2 cover metadata convention
    are written for broad reader compatibility.
    """
    target = resolve_user_path(raw_target, roots=user_accessible_roots(), must_exist=True)
    if not target.is_file() or target.suffix.lower() != ".epub":
        raise ValueError("EPUB cover insertion requires an existing .epub file.")
    if target.stat().st_size > 100 * 1024 * 1024:
        raise ValueError("Target EPUB exceeds the 100 MiB compressed-file limit.")
    _validate_epub_archive(target)
    from PIL import Image
    with Image.open(image_path) as opened:
        opened.verify()
    extension = image_path.suffix.lower()
    mime = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".webp": "image/webp"}.get(extension)
    if not mime:
        raise ValueError("EPUB cover must be PNG, JPEG or WEBP.")
    if extension == ".jpeg":
        extension = ".jpg"

    output = unique_output_path(
        target.parent,
        safe_filename(target.stem + "-with-cover", target.stem + "-with-cover", ".epub"),
    )
    temporary = output.with_name(f".{output.stem}-{uuid.uuid4().hex}.tmp.epub")
    ns = {
        "opf": "http://www.idpf.org/2007/opf",
        "dc": "http://purl.org/dc/elements/1.1/",
    }
    ET.register_namespace("", ns["opf"])
    ET.register_namespace("dc", ns["dc"])
    try:
        with zipfile.ZipFile(target, "r") as source:
            names = source.namelist()
            if len(names) != len(set(names)):
                raise ValueError("EPUB contains duplicate archive paths; refusing ambiguous edits.")
            container = ET.fromstring(source.read("META-INF/container.xml"))
            rootfile = next((node for node in container.iter() if node.tag.rsplit("}", 1)[-1] == "rootfile"), None)
            opf_path = rootfile.attrib.get("full-path", "") if rootfile is not None else ""
            if not opf_path or opf_path not in names or posixpath.normpath(opf_path).startswith("../"):
                raise ValueError("EPUB package document could not be located safely.")
            package = ET.fromstring(source.read(opf_path))
            manifest = package.find(f"{{{ns['opf']}}}manifest")
            spine = package.find(f"{{{ns['opf']}}}spine")
            metadata = package.find(f"{{{ns['opf']}}}metadata")
            if manifest is None or spine is None or metadata is None:
                raise ValueError("EPUB package is missing required metadata, manifest or spine.")
            opf_dir = posixpath.dirname(opf_path)
            image_dir = posixpath.join(opf_dir, "images") if opf_dir else "images"
            used_ids = {item.get("id", "") for item in manifest}
            image_id = "marklv-generated-cover"
            suffix = 2
            while image_id in used_ids:
                image_id = f"marklv-generated-cover-{suffix}"
                suffix += 1
            page_id = "marklv-cover-page"
            suffix = 2
            while page_id in used_ids or page_id == image_id:
                page_id = f"marklv-cover-page-{suffix}"
                suffix += 1
            image_name = f"{target.stem}-cover{extension}"
            image_archive_path = posixpath.join(image_dir, image_name)
            cover_page_name = "marklv-cover.xhtml"
            page_archive_path = posixpath.join(opf_dir, cover_page_name) if opf_dir else cover_page_name
            counter = 2
            while image_archive_path in names or page_archive_path in names:
                image_name = f"{target.stem}-cover-{counter}{extension}"
                image_archive_path = posixpath.join(image_dir, image_name)
                cover_page_name = f"marklv-cover-{counter}.xhtml"
                page_archive_path = posixpath.join(opf_dir, cover_page_name) if opf_dir else cover_page_name
                counter += 1
            image_href = posixpath.relpath(image_archive_path, opf_dir or ".")
            page_href = posixpath.relpath(page_archive_path, opf_dir or ".")
            ET.SubElement(manifest, f"{{{ns['opf']}}}item", {
                "id": image_id, "href": image_href, "media-type": mime, "properties": "cover-image"
            })
            ET.SubElement(manifest, f"{{{ns['opf']}}}item", {
                "id": page_id, "href": page_href, "media-type": "application/xhtml+xml"
            })
            spine.insert(0, ET.Element(f"{{{ns['opf']}}}itemref", {"idref": page_id, "linear": "yes"}))
            for old_meta in list(metadata):
                if old_meta.tag.rsplit("}", 1)[-1] == "meta" and old_meta.get("name") == "cover":
                    metadata.remove(old_meta)
            ET.SubElement(metadata, f"{{{ns['opf']}}}meta", {"name": "cover", "content": image_id})
            title_text = (title or target.stem).strip()[:300]
            alt = (alt_text or f"Cover artwork for {title_text}").strip()[:1000]
            xhtml = (
                '<?xml version="1.0" encoding="utf-8"?>'
                '<!DOCTYPE html><html xmlns="http://www.w3.org/1999/xhtml" '
                'xmlns:epub="http://www.idpf.org/2007/ops"><head>'
                '<title>' + html.escape(title_text) + '</title>'
                '<meta charset="utf-8"/><meta name="viewport" content="width=device-width, initial-scale=1.0"/>'
                '<style>html,body{margin:0;padding:0;text-align:center}img{display:block;max-width:100%;max-height:100vh;width:auto;height:auto;margin:0 auto}</style>'
                '</head><body><section epub:type="cover"><img src="' + html.escape(image_href, quote=True)
                + '" alt="' + html.escape(alt, quote=True) + '"/></section></body></html>'
            ).encode("utf-8")
            package_bytes = ET.tostring(package, encoding="utf-8", xml_declaration=True)
            with zipfile.ZipFile(temporary, "w") as destination:
                for info in source.infolist():
                    if info.filename == opf_path:
                        destination.writestr(info, package_bytes)
                    else:
                        destination.writestr(info, source.read(info.filename))
                destination.writestr(image_archive_path, image_path.read_bytes(), compress_type=zipfile.ZIP_DEFLATED)
                destination.writestr(page_archive_path, xhtml, compress_type=zipfile.ZIP_DEFLATED)
        with zipfile.ZipFile(temporary, "r") as check:
            if check.testzip() is not None:
                raise ValueError("Generated EPUB failed ZIP integrity validation.")
            ET.fromstring(check.read(opf_path))
            if image_archive_path not in check.namelist() or page_archive_path not in check.namelist():
                raise ValueError("Generated EPUB cover assets are missing.")
        os.replace(temporary, output)
    finally:
        temporary.unlink(missing_ok=True)
    return output


_PREVIEW_MAX_BYTES = 4 * 1024 * 1024

def _publish_preview(image_path: Path, params: dict, dashboard=None, event_loop=None, inserted_path: Path | None = None) -> bool:
    """Send a bounded image preview to authenticated dashboard WebSocket clients."""
    if dashboard is None or event_loop is None:
        return False
    try:
        image_path = resolve_user_path(image_path, roots=user_accessible_roots(), must_exist=True)
        if not image_path.is_file() or image_path.stat().st_size > _PREVIEW_MAX_BYTES:
            return False
        mime = mimetypes.guess_type(image_path.name)[0] or ""
        if mime not in {"image/png", "image/jpeg", "image/webp"}:
            return False
        payload = {
            "type": "visual_preview",
            "title": str(params.get("preview_title", "") or params.get("cover_title", "") or image_path.stem)[:120],
            "filename": image_path.name,
            "mime_type": mime,
            "data_url": f"data:{mime};base64,{base64.b64encode(image_path.read_bytes()).decode('ascii')}",
            "file_path": str(image_path),
            "inserted_path": str(inserted_path) if inserted_path else "",
        }
        future = asyncio.run_coroutine_threadsafe(dashboard.broadcast(payload), event_loop)
        future.result(timeout=4)
        return True
    except Exception:
        # Preview delivery is best-effort; never discard a saved image because a client disconnected.
        return False

def visual_studio(parameters: dict, player=None, speak=None, dashboard=None, event_loop=None) -> str:
    params = parameters or {}
    operation = str(params.get("operation", "generate") or "generate").strip().lower()
    try:
        if operation in {"generate", "edit"}:
            target = str(params.get("target_file_path", "") or "").strip()
            if target:
                _preflight_target(target, params.get("slide_index"), params.get("pptx_placement"), params.get("shape_index"), params.get("docx_insertion_mode"), params.get("insert_after_text"))
            image_path, mime = _generate(params)
            details = f"Visual saved: {image_path}\nFormat: {mime}"
            inserted = None
            target = str(params.get("target_file_path", "") or "").strip()
            if target:
                target_path = _preflight_target(target, params.get("slide_index"), params.get("pptx_placement"), params.get("shape_index"), params.get("docx_insertion_mode"), params.get("insert_after_text"))
                caption = str(params.get("caption", "") or "")
                if target_path.suffix.lower() == ".docx":
                    mode = str(params.get("docx_insertion_mode", "cover") or "cover").lower()
                    inserted = _insert_docx(image_path, target, mode, caption, str(params.get("insert_after_text", "") or ""), str(params.get("alt_text", "") or ""))
                elif target_path.suffix.lower() == ".epub":
                    inserted = _insert_epub_cover(image_path, target, str(params.get("cover_title", "") or params.get("preview_title", "") or ""), str(params.get("alt_text", "") or ""))
                else:
                    placement = str(params.get("pptx_placement", "") or ("contain" if params.get("slide_index") is not None else "append_slide"))
                    inserted = _insert_pptx(image_path, target, params.get("slide_index"), caption, placement, params.get("shape_index"), str(params.get("alt_text", "") or ""))
                details += f"\nInserted into a new copy: {inserted}"
            if _publish_preview(image_path, params, dashboard, event_loop, inserted):
                details += "\nPreview opened in the Mark-LV dashboard. Close it with the × button when finished."
            elif image_path.stat().st_size > _PREVIEW_MAX_BYTES:
                details += "\nThe image was saved, but live preview was skipped because it exceeds the 4 MiB preview limit."
            return details
        if operation == "insert":
            image_path = _image_path(str(params.get("image_path", "") or ""))
            target = str(params.get("target_file_path", "") or "").strip()
            if not target:
                raise ValueError("target_file_path is required for insert.")
            target_path = _preflight_target(target, params.get("slide_index"), params.get("pptx_placement"), params.get("shape_index"), params.get("docx_insertion_mode"), params.get("insert_after_text"))
            caption = str(params.get("caption", "") or "")
            if target_path.suffix.lower() == ".docx":
                mode = str(params.get("docx_insertion_mode", "cover") or "cover").lower()
                inserted = _insert_docx(image_path, target, mode, caption, str(params.get("insert_after_text", "") or ""), str(params.get("alt_text", "") or ""))
            elif target_path.suffix.lower() == ".epub":
                inserted = _insert_epub_cover(image_path, target, str(params.get("cover_title", "") or params.get("preview_title", "") or ""), str(params.get("alt_text", "") or ""))
            else:
                placement = str(params.get("pptx_placement", "") or ("contain" if params.get("slide_index") is not None else "append_slide"))
                inserted = _insert_pptx(image_path, target, params.get("slide_index"), caption, placement, params.get("shape_index"), str(params.get("alt_text", "") or ""))
            previewed = _publish_preview(image_path, params, dashboard, event_loop, inserted)
            result = f"Image inserted into a new copy: {inserted}\nSource image and original document/presentation were preserved."
            if previewed:
                result += "\nPreview opened in the Mark-LV dashboard."
            return result
        return "Unsupported operation. Use generate, edit, or insert."
    except ImportError as exc:
        return f"Visual Studio dependency unavailable: {exc}. Check google-genai and Pillow installation."
    except (OSError, ValueError, TypeError, KeyError, PathPolicyError, zipfile.BadZipFile) as exc:
        return f"Visual Studio failed safely: {exc}"
    except Exception as exc:
        # Avoid echoing request headers, credentials or full SDK internals into speech.
        return f"Visual generation failed: {type(exc).__name__}. Check network access, API quota and model availability."


TOOL = {
    "name": "visual_studio",
    "description": (
        "Generate or edit images with the configured Gemini image model and optionally insert the result into a new copy of an existing Word, PowerPoint or EPUB file. "
        "Operations: generate (prompt to image), edit (prompt plus existing image_path), insert (place an image into a DOCX, PPTX or EPUB). "
        "Supports covers, illustrations, posters, backgrounds and visual assets. Optional aspect_ratio and image_size tune generation. "
        "cover_title, cover_subtitle and cover_author add readable typography to a designed copy while preserving the raw generated image. "
        "Uses the app's config/api_keys.json Gemini key; never ask the user to paste the key. Existing files are preserved. "
        "DOCX supports cover or inline insertion after a matching text anchor, optional captions and accessibility alt text. PPTX supports appending a fitted slide, inserting on an existing slide, or replacing an existing picture by zero-based shape_index while preserving its bounds and stacking order. EPUB receives a cover image plus an accessible XHTML cover page, EPUB 3 cover-image metadata and EPUB 2-compatible cover metadata; the source book is preserved. "
        "The target is validated before generation to avoid wasting quota on invalid Office paths. Use exact local paths and report output paths only after success."
    ),
    "parameters": {
        "type": "OBJECT",
        "properties": {
            "operation": {"type": "STRING", "description": "generate | edit | insert"},
            "prompt": {"type": "STRING", "description": "Detailed visual brief; required for generate/edit"},
            "image_path": {"type": "STRING", "description": "Existing PNG/JPEG/WEBP source for edit or insert"},
            "target_file_path": {"type": "STRING", "description": "Optional existing DOCX/PPTX/EPUB to receive the generated image automatically; EPUB targets get a cover page and cover metadata"},
            "slide_index": {"type": "INTEGER", "description": "Zero-based PowerPoint slide index for contain/replace_picture; omit to append a new visual slide"},
            "pptx_placement": {"type": "STRING", "description": "append_slide | contain | replace_picture; replace_picture requires slide_index and shape_index"},
            "shape_index": {"type": "INTEGER", "description": "Zero-based shape index on slide_index; required to replace an existing picture"},
            "output_dir": {"type": "STRING", "description": "Optional output directory inside permitted user roots"},
            "filename": {"type": "STRING", "description": "Optional output filename, without directories"},
            "model": {"type": "STRING", "description": "Optional Gemini image model identifier; defaults to gemini-2.5-flash-image"},
            "aspect_ratio": {"type": "STRING", "description": "Optional ratio: 1:1, 2:3, 3:2, 3:4, 4:3, 4:5, 5:4, 9:16, 16:9 or 21:9"},
            "image_size": {"type": "STRING", "description": "Optional image size: 1K, 2K or 4K; model availability varies"},
            "cover_title": {"type": "STRING", "description": "Optional exact title overlaid on a designed copy"},
            "cover_subtitle": {"type": "STRING", "description": "Optional exact subtitle overlaid on a designed copy"},
            "cover_author": {"type": "STRING", "description": "Optional exact author or brand overlaid on a designed copy"},
            "docx_insertion_mode": {"type": "STRING", "description": "cover | inline; defaults to cover"},
            "insert_after_text": {"type": "STRING", "description": "Optional text fragment identifying the Word paragraph after which to insert an inline illustration"},
            "alt_text": {"type": "STRING", "description": "Optional accessible alternative text for the inserted image"},
            "caption": {"type": "STRING", "description": "Optional caption for Word or PowerPoint"}
        },
        "required": ["operation"]
    },
    "handler": visual_studio
}
