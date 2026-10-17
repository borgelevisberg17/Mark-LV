"""Tests for visual generation extraction and local Office insertion."""
from __future__ import annotations

import importlib.util
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace

from actions.visual_studio import _compose_cover_text, _extract_image, _insert_docx, _insert_pptx, _publish_preview, visual_studio


class VisualStudioUnitTests(unittest.TestCase):
    def test_extract_image_ignores_text_and_returns_image_bytes(self):
        response = SimpleNamespace(candidates=[SimpleNamespace(content=SimpleNamespace(parts=[
            SimpleNamespace(text="Here is your image."),
            SimpleNamespace(inline_data=SimpleNamespace(data=b"png-bytes", mime_type="image/png")),
        ]))])
        self.assertEqual(_extract_image(response), (b"png-bytes", "image/png"))

    def test_extract_image_fails_when_model_returns_no_image(self):
        response = SimpleNamespace(candidates=[SimpleNamespace(content=SimpleNamespace(parts=[
            SimpleNamespace(text="I cannot generate that image.")
        ]))])
        with self.assertRaisesRegex(ValueError, "no usable image"):
            _extract_image(response)

    def test_invalid_operation_is_rejected_without_network(self):
        self.assertIn("Unsupported operation", visual_studio({"operation": "delete"}))

    def test_preview_is_not_attempted_without_dashboard_context(self):
        self.assertFalse(_publish_preview(Path("missing.png"), {}, dashboard=None, event_loop=None))

    @unittest.skipUnless(importlib.util.find_spec("PIL"), "Pillow is not installed")
    def test_cover_typography_creates_separate_valid_asset(self):
        from PIL import Image

        with tempfile.TemporaryDirectory() as folder:
            raw = Path(folder) / "raw.png"
            Image.new("RGB", (640, 960), "navy").save(raw)
            designed = _compose_cover_text(raw, {
                "cover_title": "A Jornada do Conhecimento",
                "cover_subtitle": "Uma edição especial",
                "cover_author": "Borge Levisberg",
            })
            self.assertNotEqual(raw, designed)
            self.assertTrue(raw.is_file(), "Raw generated asset must be retained")
            self.assertTrue(designed.is_file())
            with Image.open(designed) as image:
                image.verify()
            self.assertNotEqual(raw.read_bytes(), designed.read_bytes())


@unittest.skipUnless(
    all(importlib.util.find_spec(name) for name in ("PIL", "docx", "pptx")),
    "Office artifact libraries are not installed",
)
class VisualStudioOfficeInsertionTests(unittest.TestCase):
    def test_insert_docx_creates_cover_copy_and_preserves_source(self):
        from PIL import Image
        from docx import Document

        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            image_path = root / "cover.png"
            Image.new("RGB", (120, 160), "white").save(image_path)
            source = root / "book.docx"
            document = Document()
            document.add_paragraph("Original content")
            document.save(source)

            output = _insert_docx(image_path, str(source))
            self.assertTrue(Path(output).is_file())
            self.assertTrue(source.is_file())
            reopened = Document(output)
            self.assertGreaterEqual(len(reopened.inline_shapes), 1)
            self.assertIn("Original content", "\n".join(p.text for p in reopened.paragraphs))

    def test_insert_docx_inline_mode_and_caption_preserve_content_order(self):
        from PIL import Image
        from docx import Document

        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            image_path = root / "illustration.png"
            Image.new("RGB", (300, 200), "white").save(image_path)
            source = root / "report.docx"
            document = Document()
            document.add_paragraph("Existing report text")
            document.save(source)

            output = _insert_docx(image_path, str(source), insertion_mode="inline", caption="Figure 1 — Illustration")
            reopened = Document(output)
            texts = [paragraph.text for paragraph in reopened.paragraphs]
            self.assertLess(texts.index("Existing report text"), texts.index("Figure 1 — Illustration"))
            self.assertTrue(source.is_file())
            self.assertGreaterEqual(len(reopened.inline_shapes), 1)

    def test_inline_illustration_can_follow_anchor_and_include_alt_text(self):
        from PIL import Image
        from docx import Document

        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            image_path = root / "diagram.png"
            Image.new("RGB", (300, 200), "white").save(image_path)
            source = root / "report.docx"
            document = Document()
            document.add_paragraph("Architecture overview")
            document.add_paragraph("This paragraph follows the anchor.")
            document.save(source)

            output = _insert_docx(
                image_path, str(source), insertion_mode="inline",
                caption="Architecture diagram", insert_after_text="Architecture overview",
                alt_text="Diagram showing the system architecture",
            )
            reopened = Document(output)
            paragraphs = reopened.paragraphs
            anchor_index = next(i for i, p in enumerate(paragraphs) if p.text == "Architecture overview")
            caption_index = next(i for i, p in enumerate(paragraphs) if p.text == "Architecture diagram")
            following_index = next(i for i, p in enumerate(paragraphs) if p.text == "This paragraph follows the anchor.")
            self.assertLess(anchor_index, caption_index)
            self.assertLess(caption_index, following_index)
            self.assertEqual(reopened.inline_shapes[0]._inline.docPr.get("descr"), "Diagram showing the system architecture")
            self.assertTrue(source.is_file())

    def test_pptx_replaces_picture_in_place_and_sets_alt_text(self):
        from PIL import Image
        from pptx import Presentation

        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            old_image = root / "old.png"
            new_image = root / "new.png"
            Image.new("RGB", (160, 90), "white").save(old_image)
            Image.new("RGB", (90, 160), "black").save(new_image)
            source = root / "deck.pptx"
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            slide.shapes.add_picture(str(old_image), 100000, 200000, width=3000000, height=1700000)
            presentation.save(source)

            output = _insert_pptx(
                new_image, str(source), slide_index=0, placement="replace_picture",
                shape_index=0, alt_text="Updated product illustration",
            )
            reopened = Presentation(output)
            shapes = reopened.slides[0].shapes
            self.assertEqual(len(shapes), 1)
            self.assertEqual(shapes[0].shape_type, 13)
            self.assertEqual(shapes[0]._element.nvPicPr.cNvPr.get("descr"), "Updated product illustration")
            self.assertTrue(source.is_file())

    def test_insert_pptx_appends_visual_slide_and_preserves_source(self):
        from PIL import Image
        from pptx import Presentation

        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            image_path = root / "poster.png"
            Image.new("RGB", (160, 90), "white").save(image_path)
            source = root / "deck.pptx"
            presentation = Presentation()
            presentation.slides.add_slide(presentation.slide_layouts[6])
            presentation.save(source)

            output = _insert_pptx(image_path, str(source), caption="Campaign poster")
            self.assertTrue(Path(output).is_file())
            self.assertTrue(source.is_file())
            reopened = Presentation(output)
            self.assertEqual(len(reopened.slides), 2)
            self.assertTrue(any(shape.shape_type == 13 for shape in reopened.slides[-1].shapes))
            self.assertTrue(any(shape.has_text_frame and shape.text == "Campaign poster" for shape in reopened.slides[-1].shapes))


if __name__ == "__main__":
    unittest.main()
