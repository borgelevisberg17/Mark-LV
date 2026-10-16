"""Integration tests for the generated Office and PDF artifacts.

Run with: pip install python-docx openpyxl python-pptx Pillow reportlab
"""
import importlib.util
import tempfile
import unittest
import zipfile
from pathlib import Path


@unittest.skipUnless(
    all(importlib.util.find_spec(name) for name in ("docx", "openpyxl", "pptx", "PIL", "reportlab")),
    "Office artifact dependencies are not installed in this test environment",
)
class OfficeArtifactIntegrationTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.output = Path(self.temp.name)

    def tearDown(self):
        self.temp.cleanup()

    def test_docx_is_editable_and_reopens(self):
        from docx import Document
        from actions.office_documents import office_documents

        result = office_documents({
            "operation": "create_docx", "title": "Quarterly report",
            "content": "Executive summary\n\nKey findings", "output_dir": str(self.output)
        })
        self.assertIn("Created editable DOCX", result)
        path = next(self.output.glob("*.docx"))
        doc = Document(path)
        self.assertEqual("Quarterly report", doc.paragraphs[0].text)
        self.assertIn("Executive summary", [p.text for p in doc.paragraphs])

    def test_epub_container_and_package_validate(self):
        from actions.office_documents import office_documents

        result = office_documents({
            "operation": "create_epub", "title": "Test ebook",
            "content": "Chapter one\n\nSecond paragraph", "output_dir": str(self.output)
        })
        self.assertIn("Created EPUB ebook", result)
        path = next(self.output.glob("*.epub"))
        with zipfile.ZipFile(path) as archive:
            self.assertEqual("mimetype", archive.namelist()[0])
            self.assertEqual("application/epub+zip", archive.read("mimetype").decode())
            self.assertIsNone(archive.testzip())
            self.assertIn(b"Test ebook", archive.read("OEBPS/chapter.xhtml"))

    def test_xlsx_is_editable_and_user_strings_are_not_formulas(self):
        from openpyxl import load_workbook
        from actions.office_spreadsheets import office_spreadsheets

        result = office_spreadsheets({
            "operation": "create_workbook", "title": "Budget",
            "headers": ["Item", "Amount"], "rows": [["Rent", 100], ["=1+1", 2]],
            "output_dir": str(self.output)
        })
        self.assertIn("Created editable XLSX", result)
        path = next(self.output.glob("*.xlsx"))
        workbook = load_workbook(path, data_only=False)
        self.assertEqual("Item", workbook.active["A1"].value)
        self.assertEqual("=1+1", workbook.active["A3"].value)
        self.assertEqual("s", workbook.active["A3"].data_type)
        workbook.close()

    def test_pptx_has_expected_slide_count_and_text(self):
        from pptx import Presentation
        from actions.office_presentations import office_presentations

        result = office_presentations({
            "operation": "create_presentation", "title": "Plan",
            "slides": [
                {"title": "Plan", "body": "Overview"},
                {"title": "Next steps", "bullets": ["Design", "Build", "Test"]}
            ],
            "output_dir": str(self.output)
        })
        self.assertIn("Created editable PPTX", result)
        path = next(self.output.glob("*.pptx"))
        presentation = Presentation(path)
        self.assertEqual(2, len(presentation.slides))
        texts = [shape.text for slide in presentation.slides for shape in slide.shapes if shape.has_text_frame]
        self.assertIn("Next steps", texts)
        self.assertIn("Build", "\n".join(texts))

    def test_image_resize_preserves_original(self):
        from PIL import Image
        from actions.image_tools import image_tools

        source = self.output / "source.png"
        Image.new("RGB", (40, 20), (20, 40, 60)).save(source)
        result = image_tools({
            "operation": "resize", "file_path": str(source), "width": 20,
            "output_dir": str(self.output)
        })
        self.assertIn("Resized image", result)
        with Image.open(source) as original:
            self.assertEqual((40, 20), original.size)
        resized = next(self.output.glob("*-resized.png"))
        with Image.open(resized) as result_image:
            self.assertEqual((20, 10), result_image.size)

    def test_text_to_pdf_creates_nonempty_pdf(self):
        from actions.document_conversion import document_conversion

        result = document_conversion({
            "operation": "text_to_pdf", "title": "PDF test",
            "content": "A paragraph for PDF output.", "output_dir": str(self.output)
        })
        self.assertIn("Created PDF", result)
        pdf = next(self.output.glob("*.pdf"))
        self.assertGreater(pdf.stat().st_size, 100)
        self.assertEqual(b"%PDF", pdf.read_bytes()[:4])


if __name__ == "__main__":
    unittest.main()
