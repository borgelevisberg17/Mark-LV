"""Regression tests for advanced Office editing and EPUB publishing."""
import importlib.util
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

from actions import office_editor


@unittest.skipUnless(all(importlib.util.find_spec(name) for name in ('docx', 'openpyxl', 'pptx', 'PIL')), 'Office editing dependencies are installed in the Office artifact test environment')
class OfficeEditorTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.docs = self.root / "Documents" / "MarkLV"
        self.docs.mkdir(parents=True)
        self.patches = [
            patch("core.path_policy.user_accessible_roots", return_value=(self.root,)),
            patch("actions.office_editor.user_accessible_roots", return_value=(self.root,)),
            patch("actions._office_common.user_accessible_roots", return_value=(self.root,)),
        ]
        for item in self.patches:
            item.start()
            self.addCleanup(item.stop)

    def test_docx_replace_preserves_source_and_run_formatting(self):
        from docx import Document

        source = self.docs / "source.docx"
        doc = Document()
        paragraph = doc.add_paragraph()
        paragraph.add_run("Aprovado: ")
        bold = paragraph.add_run("rascunho")
        bold.bold = True
        doc.save(source)

        result = office_editor.office_editor({
            "operation": "replace_text", "file_path": str(source),
            "old_text": "rascunho", "new_text": "final",
        })
        self.assertIn("Edited DOCX:", result)
        edited = self.docs / "source-edited.docx"
        self.assertTrue(edited.is_file())
        self.assertTrue(source.is_file(), "Default edit must preserve the source")
        check = Document(edited)
        self.assertEqual("Aprovado: final", check.paragraphs[0].text)
        self.assertTrue(check.paragraphs[0].runs[1].bold)

    def test_docx_adds_toc_field_and_bibliography(self):
        from docx import Document

        source = self.docs / "book.docx"
        doc = Document()
        doc.add_heading("Chapter One", level=1)
        doc.add_paragraph("Body")
        doc.save(source)
        toc_result = office_editor.office_editor({
            "operation": "add_toc", "file_path": str(source),
        })
        self.assertIn("Inserted a Word TOC field", toc_result)
        toc_copy = self.docs / "book-edited.docx"
        self.assertTrue(toc_copy.exists())
        with zipfile.ZipFile(toc_copy) as package:
            xml = package.read("word/document.xml")
            self.assertIn(b'TOC', xml)
            self.assertIn(b'fldChar', xml)
        ref_result = office_editor.office_editor({
            "operation": "add_references", "file_path": str(toc_copy),
            "references": ["Autor. (2024). Obra.", "Outra fonte. (2023)."],
        })
        self.assertIn("Added 2 bibliography entries", ref_result)
        final = self.docs / "book-edited-edited.docx"
        self.assertTrue(final.exists())
        check = Document(final)
        self.assertIn("Referências bibliográficas", [p.text for p in check.paragraphs])
        self.assertIn("Autor. (2024). Obra.", [p.text for p in check.paragraphs])

    def test_xlsx_cell_edits_are_saved_as_copy_and_strings_are_not_formulas(self):
        from openpyxl import Workbook, load_workbook

        source = self.docs / "budget.xlsx"
        workbook = Workbook()
        workbook.active.title = "Budget"
        workbook.active["A1"] = "Total"
        workbook.active["B1"] = 10
        workbook.save(source)
        result = office_editor.office_editor({
            "operation": "set_cells", "file_path": str(source), "sheet_name": "Budget",
            "cells": {"B1": 25, "A2": "=1+1"},
        })
        self.assertIn("Updated 2 cells", result)
        edited = self.docs / "budget-edited.xlsx"
        self.assertTrue(source.exists())
        check = load_workbook(edited, data_only=False)
        self.assertEqual(25, check["Budget"]["B1"].value)
        self.assertEqual("=1+1", check["Budget"]["A2"].value)
        self.assertEqual("s", check["Budget"]["A2"].data_type)
        check.close()

    def test_pptx_shape_text_edit_preserves_original_and_validates_copy(self):
        from pptx import Presentation

        source = self.docs / "deck.pptx"
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[0])
        slide.shapes.title.text = "Original title"
        deck.save(source)
        result = office_editor.office_editor({
            "operation": "replace_shape_text", "file_path": str(source),
            "slide_index": 0, "shape_index": 0, "text": "Revised title",
        })
        self.assertIn("Edited presentation:", result)
        edited = self.docs / "deck-edited.pptx"
        self.assertTrue(source.exists())
        self.assertEqual("Original title", Presentation(source).slides[0].shapes[0].text)
        self.assertEqual("Revised title", Presentation(edited).slides[0].shapes[0].text)

    def test_multichapter_epub_has_navigation_and_valid_package(self):
        result = office_editor.office_editor({
            "operation": "create_epub", "title": "Livro de teste",
            "output_dir": str(self.docs),
            "chapters": [
                {"title": "Introdução", "content": "Primeiro parágrafo.\n\nSegundo parágrafo."},
                {"title": "Conclusão", "content": "Texto final."},
            ],
        })
        self.assertIn("Chapters: 2", result)
        path = self.docs / "Livro de teste.epub"
        self.assertTrue(path.exists())
        with zipfile.ZipFile(path) as book:
            self.assertEqual("mimetype", book.namelist()[0])
            self.assertIsNone(book.testzip())
            self.assertIn(b"Introdu\xc3\xa7\xc3\xa3o", book.read("OEBPS/nav.xhtml"))
            self.assertIn("OEBPS/chapter-002.xhtml", book.namelist())

    def test_action_is_discoverable_by_existing_registry(self):
        from core.action_loader import discover_actions

        registry = discover_actions(Path(__file__).resolve().parents[1] / "actions", logger=lambda _: None)
        self.assertTrue(registry.has("office_editor"))


if __name__ == "__main__":
    unittest.main()
